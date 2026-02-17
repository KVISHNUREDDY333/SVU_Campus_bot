import markdown
from xhtml2pdf import pisa
from io import BytesIO
from docx import Document
from docx.shared import Pt
import re

def generate_pdf(markdown_text: str) -> BytesIO:
    """
    Converts Markdown text to a PDF file in memory.
    """
    # Convert Markdown to HTML
    html_content = markdown.markdown(markdown_text)
    
    # Add basic styling for the PDF
    styled_html = f"""
    <html>
    <head>
        <style>
            body {{ font-family: Helvetica, sans-serif; font-size: 12pt; line-height: 1.5; }}
            h1 {{ font-size: 18pt; color: #2c3e50; border-bottom: 2px solid #2c3e50; padding-bottom: 5px; margin-top: 20px; }}
            h2 {{ font-size: 16pt; color: #34495e; margin-top: 15px; margin-bottom: 10px; }}
            h3 {{ font-size: 14pt; color: #7f8c8d; margin-top: 10px; }}
            p {{ margin-bottom: 10px; text-align: justify; }}
            ul {{ margin-bottom: 10px; }}
            li {{ margin-bottom: 5px; }}
            strong {{ font-weight: bold; color: #2c3e50; }}
        </style>
    </head>
    <body>
        {html_content}
    </body>
    </html>
    """
    
    pdf_buffer = BytesIO()
    pisa_status = pisa.CreatePDF(styled_html, dest=pdf_buffer)
    
    if pisa_status.err:
        raise Exception("PDF generation failed")
        
    pdf_buffer.seek(0)
    return pdf_buffer

def generate_docx(markdown_text: str) -> BytesIO:
    """
    Converts Markdown text to a Word (.docx) file in memory.
    """
    document = Document()
    
    # Basic Markdown parsing (Note: This is a simplified parser)
    lines = markdown_text.split('\n')
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('# '):
            document.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            document.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            document.add_heading(line[4:], level=3)
        elif line.startswith('- ') or line.startswith('* '):
            # Remove bold/italic markers for cleaner text
            clean_text = line[2:].replace('**', '').replace('*', '')
            document.add_paragraph(clean_text, style='List Bullet')
        else:
            # Regular paragraph
             # Remove bold/italic markers for cleaner text
            clean_text = line.replace('**', '').replace('*', '')
            document.add_paragraph(clean_text)

    docx_buffer = BytesIO()
    document.save(docx_buffer)
    docx_buffer.seek(0)
    return docx_buffer
